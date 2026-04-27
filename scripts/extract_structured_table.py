from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def _detect_header_row(df: pd.DataFrame, max_rows: int = 15) -> int:
    best_idx = 0
    best_score = -1
    upper = min(max_rows, len(df))
    for i in range(upper):
        vals = [
            str(x).strip()
            for x in df.iloc[i].tolist()
            if str(x).strip() and str(x).strip() != "nan"
        ]
        score = len(vals)
        if score > best_score:
            best_score = score
            best_idx = i
    return best_idx


def _clean_cell(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


def _dedup_columns(cols: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for c in cols:
        base = c or "COL"
        n = seen.get(base, 0)
        if n == 0:
            out.append(base)
        else:
            out.append(f"{base}_{n + 1}")
        seen[base] = n + 1
    return out


def _forward_fill(values: list[str]) -> list[str]:
    out: list[str] = []
    last = ""
    for v in values:
        t = str(v or "").strip()
        if t and t.lower() != "nan":
            last = t
        out.append(last if last else t)
    return out


def _extract_multi_header_table(df: pd.DataFrame) -> pd.DataFrame | None:
    if len(df) < 4:
        return None
    r1 = [_clean_cell(v) for v in df.iloc[1].tolist()]
    r2 = [_clean_cell(v) for v in df.iloc[2].tolist()]
    if "机型" not in r1 or "公司" not in r1:
        return None
    if not any(x in " ".join(r2) for x in ("环比", "同比", "月", "日")):
        return None
    r1f = _forward_fill(r1)
    r2f = _forward_fill(r2)
    cols: list[str] = []
    for a, b in zip(r1f, r2f):
        if a in ("机型", "公司"):
            cols.append(a)
            continue
        bb = b if b not in ("环比", "同比") else b
        if a and bb:
            cols.append(f"{a}_{bb}")
        else:
            cols.append(a or bb or "")
    cols = _dedup_columns([c if c else "COL" for c in cols])
    data = df.iloc[3:].copy()
    data = data.iloc[:, : len(cols)]
    data.columns = cols
    data = data.fillna("")
    data = data[data.apply(lambda r: any(str(x).strip() for x in r.tolist()), axis=1)]
    return data.reset_index(drop=True)


def _openpyxl_sheet_to_rows(path: Path) -> tuple[list[dict], int]:
    wb = load_workbook(path, data_only=True)
    sheet_rows: list[list[object]] = []
    try:
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
            if not rows:
                continue
            header_idx = 0
            best_score = -1
            for idx, row in enumerate(rows[:15]):
                score = len([v for v in row if _clean_cell(v)])
                if score > best_score:
                    best_score = score
                    header_idx = idx
            headers = [_clean_cell(v) for v in rows[header_idx]]
            headers = [h for h in headers if h]
            if not headers:
                continue
            deduped = _dedup_columns(headers)
            for row in rows[header_idx + 1 :]:
                values = [_clean_cell(v) for v in row[: len(deduped)]]
                if not any(values):
                    continue
                if len(values) < len(deduped):
                    values.extend([""] * (len(deduped) - len(values)))
                sheet_rows.append(dict(zip(deduped, values)))
        return sheet_rows, len(wb.sheetnames)
    finally:
        wb.close()


def _is_meaningful_row(row: dict) -> bool:
    id_cols = ["航班号", "航班日期", "航段"]
    if any(str(row.get(c, "")).strip() for c in id_cols):
        return True

    values = [str(v).strip() for v in row.values() if str(v).strip()]
    if len(values) >= 3:
        return True

    noise_tokens = ("备注", "数据最后更新时间", "新开航线")
    text = " ".join(values)
    if any(t in text for t in noise_tokens):
        return False
    return len(values) >= 2


def extract_table(path: Path) -> dict:
    if not hasattr(pd, "ExcelFile"):
        rows, sheet_count = _openpyxl_sheet_to_rows(path)
        meaningful_rows = [r for r in rows if _is_meaningful_row(r)]
        return {
            "columns": list(rows[0].keys()) if rows else [],
            "rows": rows,
            "meaningful_row_count": len(meaningful_rows),
            "sheet_count": sheet_count,
        }

    xls = pd.ExcelFile(path)
    sheet_frames: list[pd.DataFrame] = []
    for sheet in xls.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet, header=None)
        if df.empty:
            continue
        multi = _extract_multi_header_table(df)
        if multi is not None and not multi.empty:
            sheet_frames.append(multi)
            continue
        header_idx = _detect_header_row(df)
        headers = [_clean_cell(v) for v in df.iloc[header_idx].tolist()]
        headers = [h for h in headers if h]
        if not headers:
            continue
        data = df.iloc[header_idx + 1 :].copy()
        data = data.iloc[:, : len(headers)]
        data.columns = _dedup_columns(headers)
        data = data.fillna("")
        # Drop rows with all-empty values.
        data = data[
            data.apply(lambda r: any(str(x).strip() for x in r.tolist()), axis=1)
        ]
        data = data.reset_index(drop=True)
        if not data.empty:
            sheet_frames.append(data)

    if not sheet_frames:
        return {"columns": [], "rows": [], "sheet_count": len(xls.sheet_names)}

    if len(sheet_frames) == 1:
        merged = sheet_frames[0]
    else:
        # FineReport exports often split wide tables across Page1/Page2/Page3.
        max_len = max(len(f) for f in sheet_frames)
        padded = []
        for f in sheet_frames:
            if len(f) < max_len:
                pad = pd.DataFrame(
                    [[""] * len(f.columns)] * (max_len - len(f)), columns=f.columns
                )
                f = pd.concat([f, pad], ignore_index=True)
            padded.append(f.reset_index(drop=True))
        merged = pd.concat(padded, axis=1)
        merged.columns = _dedup_columns([str(c) for c in merged.columns])

    rows = merged.to_dict(orient="records")
    meaningful_rows = [r for r in rows if _is_meaningful_row(r)]
    return {
        "columns": list(merged.columns),
        "rows": rows,
        "meaningful_row_count": len(meaningful_rows),
        "sheet_count": len(xls.sheet_names),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract structured table rows from report excel."
    )
    parser.add_argument("file_path")
    args = parser.parse_args()
    result = extract_table(Path(args.file_path))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
