from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

AIRLINE_NAMES = {
    "航空股份",
    "首都航空",
    "天津航空",
    "祥鹏航空",
    "西部航空",
    "北部湾航空",
    "福州航空",
    "乌鲁木齐航空",
    "长安航空",
    "金鹏航空",
    "香港航空",
    "加纳航空",
}


def extract_adjusted_profit_overview_rows(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), values_only=True))
    if len(rows) < 5:
        return []

    top = rows[1]
    sub = rows[2]
    headers: list[str] = []
    current_group = ""
    for idx in range(len(sub)):
        g = str(top[idx] or "").strip()
        s = str(sub[idx] or "").strip()
        if g:
            current_group = g
        if current_group and s:
            headers.append(f"{current_group}_{s}")
        elif current_group:
            headers.append(current_group)
        else:
            headers.append(s)

    airline_col = 0
    yoy_col = None
    rank_col = None
    for i, h in enumerate(headers):
        hs = str(h)
        if ("净利润" in hs) and ("同比" in hs):
            yoy_col = i
        if ("净利润" in hs) and ("本期" in hs):
            airline_col = 0
        if ("排名" in hs) and rank_col is None:
            rank_col = i
    if yoy_col is None:
        return []

    out: list[dict] = []
    for row in rows[4:]:
        airline = str(row[0] or "").strip()
        if (not airline) or (airline not in AIRLINE_NAMES):
            continue
        yoy = row[yoy_col] if yoy_col < len(row) else None
        rank = row[rank_col] if rank_col is not None and rank_col < len(row) else None
        try:
            yoy_v = float(yoy)
        except Exception:
            continue
        item = {
            "航司": airline,
            "净利润_同比": yoy_v * 100.0 if abs(yoy_v) <= 10 else yoy_v,
        }
        if rank is not None and str(rank).strip():
            item["排名"] = rank
        out.append(item)
    return out
