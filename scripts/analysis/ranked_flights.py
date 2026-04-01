from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


def _normalize_route_text(route_text: object) -> str:
    s = str(route_text or "").strip()
    s = s.replace("→", "-").replace("—", "-").replace("–", "-").replace("=", "-")
    s = re.sub(r"\s+", "", s)
    return s


def _extract_route_pairs(route_text: str) -> list[str]:
    s = _normalize_route_text(route_text)
    if not s:
        return []
    if not re.search(r"[\u4e00-\u9fff]", s):
        return []
    parts = [x for x in re.split(r"[=-]+", s) if x]
    if len(parts) < 2:
        return []
    out: list[str] = []
    for i in range(len(parts) - 1):
        a = parts[i]
        b = parts[i + 1]
        if not re.search(r"[\u4e00-\u9fff]", a) or not re.search(r"[\u4e00-\u9fff]", b):
            continue
        out.append(f"{a}-{b}")
    return out


def analyze_first_flight_bottom10(intent: dict, extracted: dict, source_meta: dict) -> dict:
    rows = extracted.get("rows") or []
    jd_hits: list[str] = []
    seen_jd = set()
    for r in rows:
        comp2 = str(r.get("公司_2", "")).strip()
        flt2 = str(r.get("往返航班号_2", "")).strip().upper()
        route_raw = str(r.get("机型_2", "")).strip() or str(r.get("航段班次↑↓_2", "")).strip()
        if not route_raw:
            continue
        if ("首都航空" in comp2) or ("首都航空" in flt2) or flt2.startswith("JD"):
            pairs = _extract_route_pairs(route_raw)
            if not pairs:
                continue
            rp = pairs[0]
            if rp not in seen_jd:
                seen_jd.add(rp)
                jd_hits.append(rp)
    if not jd_hits:
        route_cols: list[str] = []
        if rows:
            for k in rows[0].keys():
                ks = str(k or "")
                if ("航段班次" in ks) or ("航线" in ks):
                    route_cols.append(ks)
        seen = set()
        for r in rows:
            for c in route_cols:
                for rp in _extract_route_pairs(r.get(c, "")):
                    if rp not in seen:
                        seen.add(rp)
                        jd_hits.append(rp)
    return {
        "ok": True,
        "analysis_engine": "ranked_flights",
        "matched_route": None,
        "matched_dates": [],
        "issues": [],
        "advice": [],
        "summary": f"后十首航航班共 {len(jd_hits)} 条",
        "first_flight_routes": jd_hits,
    }


def render_first_flight_bottom10_answer(analysis_result: dict, source_meta: dict) -> str:
    report_name = str(source_meta.get("report_name") or "")
    source_path = str(source_meta.get("file_path") or "")
    routes = [str(x) for x in (analysis_result.get("first_flight_routes") or []) if str(x).strip()]
    return (
        f"命中报表: {report_name}\n"
        f"来源: {source_path}\n"
        f"后十首航航班共 {len(routes)} 条\n"
        f"分别是: {'、'.join(routes) if routes else '-'}"
    )


def analyze_top_metric_flight(intent: dict, source_meta: dict) -> dict:
    source_path = str(source_meta.get("file_path") or "")
    metric_text = str((intent or {}).get("metric") or "").strip()
    report_name = str(source_meta.get("report_name") or "")
    if metric_text not in {"小时边际贡献", "总边贡"}:
        return {
            "ok": False,
            "analysis_engine": "ranked_flights",
            "matched_route": None,
            "matched_dates": [],
            "issues": [],
            "advice": [],
            "summary": f"暂不支持指标: {metric_text}",
        }
    metric_col = "小时边际贡献(万元)↑↓" if metric_text == "小时边际贡献" else "总边贡(万元)↑↓"
    try:
        wb = load_workbook(source_path, data_only=True)
        ws = wb[wb.sheetnames[0]]
    except Exception:
        return {
            "ok": False,
            "analysis_engine": "ranked_flights",
            "matched_route": None,
            "matched_dates": [],
            "issues": [],
            "advice": [],
            "summary": f"无法读取报表文件。报表: {report_name} 来源: {source_path}",
        }
    header_map: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(2, c).value
        if v is None:
            continue
        header_map[str(v).strip()] = c
    flight_col = header_map.get("往返航班号")
    route_col = header_map.get("往返航线")
    company_col = header_map.get("公司")
    target_col = header_map.get(metric_col)
    if not flight_col or not target_col:
        return {
            "ok": False,
            "analysis_engine": "ranked_flights",
            "matched_route": None,
            "matched_dates": [],
            "issues": [],
            "advice": [],
            "summary": f"未在报表中识别到 {metric_text} 列。",
        }

    def _to_float(v: object) -> float | None:
        s = str(v or "").strip()
        if not s:
            return None
        m = re.search(r"-?\d+(?:\.\d+)?", s.replace(",", ""))
        if not m:
            return None
        try:
            return float(m.group(0))
        except Exception:
            return None

    best: dict | None = None
    for r in range(3, ws.max_row + 1):
        flight_no = str(ws.cell(r, flight_col).value or "").strip()
        if not flight_no:
            continue
        value = _to_float(ws.cell(r, target_col).value)
        if value is None:
            continue
        item = {
            "rank": int(_to_float(ws.cell(r, 2).value) or 0),
            "company": str(ws.cell(r, company_col).value or "").strip() if company_col else "",
            "flight_no": flight_no,
            "route": str(ws.cell(r, route_col).value or "").strip() if route_col else "",
            "value": value,
        }
        if best is None or item["value"] > best["value"]:
            best = item
    if not best:
        return {
            "ok": False,
            "analysis_engine": "ranked_flights",
            "matched_route": None,
            "matched_dates": [],
            "issues": [],
            "advice": [],
            "summary": "未在报表中识别到有效航班数据。",
        }
    return {
        "ok": True,
        "analysis_engine": "ranked_flights",
        "matched_route": None,
        "matched_dates": [],
        "issues": [],
        "advice": [],
        "summary": f"{metric_text}最高的航班: {best['flight_no']}",
        "best_item": best,
    }


def render_top_metric_flight_answer(analysis_result: dict, intent: dict, source_meta: dict) -> str:
    report_name = str(source_meta.get("report_name") or "")
    source_path = str(source_meta.get("file_path") or "")
    metric_text = str((intent or {}).get("metric") or "").strip()
    if not analysis_result.get("ok"):
        return f"报表: {report_name}\n来源: {source_path}\n{str(analysis_result.get('summary') or '未识别到有效结果。')}"
    best = analysis_result.get("best_item") or {}
    metric_label = "小时边际贡献" if metric_text == "小时边际贡献" else "总边贡"
    return (
        f"命中报表: {report_name}\n"
        f"来源: {source_path}\n"
        f"{metric_label}最高的航班: {best.get('flight_no') or '-'}\n"
        f"航线: {best.get('route') or '-'}\n"
        f"公司: {best.get('company') or '-'}\n"
        f"{metric_label}: {float(best.get('value') or 0.0):.1f} 万元"
    )
