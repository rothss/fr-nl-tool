"""
Excel export snapshot extraction.

Reads an exported xlsx file and produces a standardized export_snapshot
data structure suitable for consistency comparison.
"""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .normalize_table import fill_merged_cells, normalize_rows


def _clean_cell(v: object) -> str:
    """Clean cell value for header detection."""
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


def _sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _file_hash(path: Path) -> str:
    """Compute SHA-256 hash of a file."""
    return _sha256_hex(path.read_bytes())


def _dedup_columns(cols: list[str]) -> list[str]:
    """Deduplicate column names by appending _2, _3, etc."""
    seen: dict[str, int] = {}
    out: list[str] = []
    for c in cols:
        base = c or "COL"
        n = seen.get(base, 0)
        if n == 0:
            out.append(base)
        else:
            out.append(f"{base}_{n + 1}")
        seen[base] = n + 1
    return out


def _detect_header_row(rows: list[list[object]], max_rows: int = 15) -> int:
    """Detect the header row index in an Excel sheet."""
    best_idx = 0
    best_score = -1
    upper = min(max_rows, len(rows))
    for i in range(upper):
        score = len([v for v in rows[i] if _clean_cell(v)])
        if score > best_score:
            best_score = score
            best_idx = i
    return best_idx


def _is_meaningful_row(row: dict) -> bool:
    """Check if a row contains meaningful data (not just noise/meta)."""
    noise_tokens = ("备注", "数据最后更新时间", "新开航线", "制表", "审核")
    text = " ".join(str(v).strip() for v in row.values() if str(v).strip())
    if any(t in text for t in noise_tokens):
        return False
    values = [str(v).strip() for v in row.values() if str(v).strip()]
    return len(values) >= 1


def extract_export_snapshot(
    file_path: Path,
    compare_config: dict | None = None,
) -> dict[str, Any]:
    """Extract a standardized export snapshot from an Excel file.

    Args:
        file_path: Path to the exported xlsx file
        compare_config: Comparison configuration (for normalization hints)

    Returns:
        Export snapshot dict with structure:
        {
            "source": "export",
            "file": str,
            "file_hash": str,
            "workbook": {"sheet_name": str, "sheet_count": int},
            "columns": [...],
            "rows": [{...}, ...],
            "row_count": int,
            "meaningful_row_count": int,
            "parsed_at": str,
        }
    """
    cfg = compare_config or {}
    key_columns = cfg.get("key_columns", [])
    value_columns = cfg.get("value_columns", [])
    date_columns = cfg.get("date_columns", [])
    percent_columns = cfg.get("percent_columns", [])
    numeric_columns = cfg.get("numeric_columns", [])
    ignore_columns = cfg.get("ignore_columns", [])
    aliases = cfg.get("aliases", {})
    fill_merged = cfg.get("fill_merged_cells", True)
    allow_extra_export_rows = cfg.get("allow_extra_export_rows", True)

    # Config-driven extraction hints
    target_sheet_name = cfg.get("sheet_name")
    explicit_header_row = cfg.get("header_row")      # 1-based row index in config
    data_start_row = cfg.get("data_start_row")        # 1-based row index in config

    wb = load_workbook(file_path, data_only=True)
    try:
        sheet_names = wb.sheetnames
        all_rows: list[dict] = []
        # Use configured sheet_name, filter to first match, or fall back to all
        sheets_to_read = [target_sheet_name] if target_sheet_name and target_sheet_name in sheet_names else sheet_names
        primary_sheet_name = sheets_to_read[0] if sheets_to_read else "Sheet1"

        for sheet_name in sheets_to_read:
            ws = wb[sheet_name]
            raw_rows = [
                [cell for cell in row]
                for row in ws.iter_rows(values_only=True)
            ]
            if not raw_rows:
                continue

            # ── Header detection (config overrides auto-detect) ──
            if explicit_header_row and explicit_header_row >= 1:
                header_idx = explicit_header_row - 1  # 1-based → 0-based
            else:
                header_idx = _detect_header_row(raw_rows)
            headers = [_clean_cell(v) for v in raw_rows[header_idx]]
            headers = [h for h in headers if h]
            if not headers:
                continue
            headers = _dedup_columns(headers)

            # ── Data row range ──
            start_idx = header_idx + 1
            if data_start_row and data_start_row >= 1:
                start_idx = data_start_row - 1  # 1-based → 0-based
            row_payload = raw_rows[start_idx:]

            # ── Fill merged cells if requested ──
            if fill_merged:
                row_payload = fill_merged_cells(row_payload)

            for row_data in row_payload:
                values = [_clean_cell(v) for v in row_data[:len(headers)]]
                if len(values) < len(headers):
                    values.extend([""] * (len(headers) - len(values)))
                row_dict = dict(zip(headers, values))
                if _is_meaningful_row(row_dict):
                    all_rows.append(row_dict)

        # Determine which columns to use for comparison
        all_columns = list(all_rows[0].keys()) if all_rows else []
        identified_value_cols = value_columns or all_columns

        # Normalize
        normalized_rows = normalize_rows(
            all_rows,
            columns=all_columns,
            date_columns=date_columns,
            percent_columns=percent_columns,
            numeric_columns=numeric_columns,
            fill_merged=fill_merged,
            aliases=aliases,
            ignore_columns=ignore_columns,
        )

        snapshot = {
            "source": "export",
            "file": str(file_path),
            "file_hash": _file_hash(file_path),
            "workbook": {
                "sheet_name": primary_sheet_name,
                "sheet_count": len(sheet_names),
            },
            "columns": all_columns,
            "rows": normalized_rows,
            "row_count": len(normalized_rows),
            "meaningful_row_count": len(all_rows),
            "parsed_at": datetime.now(timezone.utc).isoformat(),
        }
        return snapshot

    finally:
        wb.close()


def save_export_snapshot(snapshot: dict, output_path: Path) -> None:
    """Save export snapshot as JSON to disk."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
