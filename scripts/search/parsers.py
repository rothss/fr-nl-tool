from __future__ import annotations

from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .normalizer import normalize_value


def _col_index(ref: str) -> int:
    v = 0
    for ch in ref:
        if ch.isalpha():
            v = v * 26 + (ord(ch.upper()) - ord("A") + 1)
    return v


def _parse_shared_strings(zf: ZipFile) -> list[str]:
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    try:
        root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    result: list[str] = []
    for si in root.findall("a:si", ns):
        parts = [node.text or "" for node in si.iter() if node.tag.endswith("}t") and node.text]
        result.append("".join(parts))
    return result


def _parse_sheet_targets(zf: ZipFile) -> list[tuple[str, str]]:
    ns_m = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    ns_r = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}

    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rel_map: dict[str, str] = {}
    for rel in rels:
        rid = rel.attrib.get("Id")
        tgt = rel.attrib.get("Target")
        if rid and tgt:
            rel_map[rid] = tgt.replace("\\", "/")

    sheets: list[tuple[str, str]] = []
    for sh in wb_root.findall("a:sheets/a:sheet", {**ns_m, **ns_r}):
        name = sh.attrib.get("name", "")
        rid = sh.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        if not rid or rid not in rel_map:
            continue
        tgt = rel_map[rid]
        if not tgt.startswith("/"):
            tgt = f"xl/{tgt.lstrip('/')}"
        sheets.append((name, tgt))
    return sheets


def _parse_xlsx_sheet_xml(
    zf: ZipFile, sheet_path: str, shared_strings: list[str], drop_numeric: bool
) -> list[tuple[int, int, str]]:
    """XML 回退解析单个 Sheet，含合并单元格回填"""
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(zf.read(sheet_path))

    cells: dict[tuple[int, int], str] = {}
    sd = root.find("a:sheetData", ns)
    if sd is not None:
        for row in sd.findall("a:row", ns):
            for cell in row.findall("a:c", ns):
                ref = cell.attrib.get("r", "")
                col_ref = "".join(ch for ch in ref if ch.isalpha())
                row_ref = "".join(ch for ch in ref if ch.isdigit())
                if not col_ref or not row_ref:
                    continue
                rn, cn = int(row_ref), _col_index(col_ref)

                ct = cell.attrib.get("t")
                val: str | None = None
                if ct == "inlineStr":
                    is_el = cell.find("a:is", ns)
                    if is_el is not None:
                        parts = [n.text or "" for n in is_el.iter() if n.tag.endswith("}t") and n.text]
                        val = "".join(parts)
                else:
                    formula = cell.find("a:f", ns)
                    v_el = cell.find("a:v", ns)
                    if formula is not None and formula.text:
                        val = formula.text
                    elif v_el is not None and v_el.text is not None:
                        raw = v_el.text
                        if ct == "s":
                            try:
                                val = shared_strings[int(raw)]
                            except Exception:
                                val = raw
                        else:
                            val = raw
                nv = normalize_value(val, drop_numeric=drop_numeric)
                if nv is not None:
                    cells[(rn, cn)] = nv

    mc = root.find("a:mergeCells", ns)
    if mc is not None:
        for m in mc.findall("a:mergeCell", ns):
            mref = m.attrib.get("ref")
            if not mref:
                continue
            min_col, min_row, max_c, max_r = range_boundaries(mref)
            tl = cells.get((min_row, min_col))
            if tl is None:
                continue
            for row_no in range(min_row, max_r + 1):
                for col_no in range(min_col, max_c + 1):
                    cells.setdefault((row_no, col_no), tl)

    return [(r, c, v) for (r, c), v in sorted(cells.items())]


def parse_xlsx(path: Path, drop_numeric: bool = False) -> dict[str, list[tuple[int, int, str]]]:
    """openpyxl 主解析，失败回退 XML"""
    try:
        wb = load_workbook(filename=path, read_only=True, data_only=False)
        result: dict[str, list[tuple[int, int, str]]] = {}
        try:
            for name in wb.sheetnames:
                ws = wb[name]
                try:
                    ws.reset_dimensions()
                except Exception:
                    pass
                rows: list[tuple[int, int, str]] = []
                for row in ws.iter_rows():
                    for cell in row:
                        nv = normalize_value(cell.value, drop_numeric=drop_numeric)
                        if nv is not None:
                            rows.append((cell.row, cell.column, nv))
                result[name] = rows
        finally:
            wb.close()
        return result
    except Exception:
        return _parse_xlsx_fallback(path, drop_numeric)


def _parse_xlsx_fallback(path: Path, drop_numeric: bool) -> dict[str, list[tuple[int, int, str]]]:
    result: dict[str, list[tuple[int, int, str]]] = {}
    with ZipFile(path) as zf:
        ss = _parse_shared_strings(zf)
        for sheet_name, sheet_path in _parse_sheet_targets(zf):
            result[sheet_name] = _parse_xlsx_sheet_xml(zf, sheet_path, ss, drop_numeric)
    return result


def parse_xls(path: Path, drop_numeric: bool = False) -> dict[str, list[tuple[int, int, str]]]:
    result: dict[str, list[tuple[int, int, str]]] = {}
    xls = pd.ExcelFile(path, engine="xlrd")
    for name in xls.sheet_names:
        df: Any = pd.read_excel(path, sheet_name=name, header=None, dtype=object, engine="xlrd")
        max_r, max_c = int(df.shape[0]), int(df.shape[1])
        rows: list[tuple[int, int, str]] = []
        for ri in range(max_r):
            for ci in range(max_c):
                nv = normalize_value(df.iat[ri, ci], drop_numeric=drop_numeric)
                if nv is not None:
                    rows.append((ri + 1, ci + 1, nv))
        result[name] = rows
    return result


def parse_csv(path: Path, drop_numeric: bool = False) -> dict[str, list[tuple[int, int, str]]]:
    df = pd.read_csv(path, header=None, dtype=object, keep_default_na=False, encoding_errors="ignore")
    max_r, max_c = int(df.shape[0]), int(df.shape[1])
    rows: list[tuple[int, int, str]] = []
    for ri in range(max_r):
        for ci in range(max_c):
            nv = normalize_value(df.iat[ri, ci], drop_numeric=drop_numeric)
            if nv is not None:
                rows.append((ri + 1, ci + 1, nv))
    return {"CSV": rows}


SUPPORTED = {".xlsx", ".xlsm", ".xls", ".csv"}


def parse_file(path: Path, drop_numeric: bool = False) -> dict[str, list[tuple[int, int, str]]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return parse_xlsx(path, drop_numeric=drop_numeric)
    if suffix == ".xls":
        return parse_xls(path, drop_numeric=drop_numeric)
    if suffix == ".csv":
        return parse_csv(path, drop_numeric=drop_numeric)
    raise ValueError(f"Unsupported file type: {path}")
