from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from common import default_catalog_db, default_mirror_root

SUPPORTED = {".xlsx", ".xlsm", ".xls"}


def find_header_rows(path: Path, max_scan_rows: int = 15) -> list[tuple[str, int, str]]:
    out: list[tuple[str, int, str]] = []
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            best_row = 1
            best_score = -1
            best_headers: list[str] = []
            upper = min(max_scan_rows, ws.max_row or max_scan_rows)
            for r in range(1, upper + 1):
                vals = []
                for c in range(1, min((ws.max_column or 1), 40) + 1):
                    v = ws.cell(r, c).value
                    txt = str(v).strip() if v is not None else ""
                    if txt:
                        vals.append(txt)
                if not vals:
                    continue
                score = len(vals)
                if score > best_score:
                    best_score = score
                    best_row = r
                    best_headers = vals
            out.append((sheet_name, best_row, " | ".join(best_headers)))
    finally:
        wb.close()
    return out


def init_db(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        DROP TABLE IF EXISTS reports_fts;
        DROP TABLE IF EXISTS report_headers;
        DROP TABLE IF EXISTS reports;

        CREATE TABLE reports (
            id INTEGER PRIMARY KEY,
            report_name TEXT NOT NULL,
            file_path TEXT NOT NULL UNIQUE,
            dir_path TEXT NOT NULL,
            mtime REAL NOT NULL,
            size_bytes INTEGER NOT NULL
        );

        CREATE TABLE report_headers (
            id INTEGER PRIMARY KEY,
            report_id INTEGER NOT NULL,
            sheet_name TEXT NOT NULL,
            header_row INTEGER NOT NULL,
            header_text TEXT NOT NULL,
            FOREIGN KEY (report_id) REFERENCES reports(id)
        );

        CREATE VIRTUAL TABLE reports_fts USING fts5(
            report_name,
            dir_path,
            headers_text,
            file_path UNINDEXED,
            report_id UNINDEXED
        );
        """
    )
    conn.commit()


def build_catalog(root: Path, db_path: Path) -> dict[str, int]:
    files = [
        p
        for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith("~$")
    ]
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    stats = {"reports": 0, "headers": 0, "failed": 0}
    try:
        init_db(conn)
        for path in files:
            try:
                rel = path.relative_to(root)
                st = path.stat()
                cur = conn.execute(
                    """
                    INSERT INTO reports(report_name, file_path, dir_path, mtime, size_bytes)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        path.stem,
                        str(path),
                        str(rel.parent).replace("\\", "/"),
                        st.st_mtime,
                        st.st_size,
                    ),
                )
                report_id = int(cur.lastrowid)
                headers = find_header_rows(path)
                headers_text = []
                for sheet_name, row_no, header_text in headers:
                    conn.execute(
                        """
                        INSERT INTO report_headers(report_id, sheet_name, header_row, header_text)
                        VALUES (?, ?, ?, ?)
                        """,
                        (report_id, sheet_name, row_no, header_text),
                    )
                    headers_text.append(header_text)
                    stats["headers"] += 1
                conn.execute(
                    """
                    INSERT INTO reports_fts(report_name, dir_path, headers_text, file_path, report_id)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        path.stem,
                        str(rel.parent).replace("\\", "/"),
                        " || ".join(headers_text),
                        str(path),
                        report_id,
                    ),
                )
                stats["reports"] += 1
            except Exception:
                stats["failed"] += 1
                conn.rollback()
            else:
                conn.commit()
    finally:
        conn.close()
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="Build local OPM report catalog database.")
    parser.add_argument("--root", default=str(default_mirror_root()))
    parser.add_argument("--db", default=str(default_catalog_db()))
    args = parser.parse_args()

    stats = build_catalog(Path(args.root), Path(args.db))
    print(f"Catalog built: reports={stats['reports']} headers={stats['headers']} failed={stats['failed']}")
    print(f"Database: {args.db}")


if __name__ == "__main__":
    main()

